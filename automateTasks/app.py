import openpyxl as xl
from openpyxl.chart import BarChart, Reference


# Let's say we have a transaction data sheet, but the price in 3rd column is calculated incorrectly.
# We need to correct it by getting 90% of those values.
# But making this change manually for thousands of rows is a time-consuming task
# Automations come to play in this kind of situations

workbook = xl.load_workbook('transactions .xlsx')
sheet = workbook['Sheet1']

# We can get the cells from both of the following ways
cell = sheet['A1']
cell2 = sheet.cell(row=1, column=1)
print(cell.value)
print(cell2.value)

#Get the rows number
rowsCount = sheet.max_row
#print(rowsCount)

for row in range(2, rowsCount+1):
    cell = sheet.cell(row=row, column=3)
    corrected_price = cell.value*0.9
    print(cell.value, corrected_price)
    corrected_price_cell = sheet.cell(row=row, column=4)
    corrected_price_cell.value = corrected_price

#Adding a chart to xlsx sheet
chart_values = Reference(
    sheet,
    min_row=2,
    max_row=rowsCount,
    min_col=4,
    max_col=4
)

chart = BarChart()
chart.add_data(chart_values)
sheet.add_chart(chart)

workbook.save('transactions2.xlsx')