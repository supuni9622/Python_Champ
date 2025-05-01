import openpyxl as xl
from openpyxl.chart import BarChart, Reference

def process_files(input_file, output_file, sheet_name):
    workbook = xl.load_workbook(input_file)
    sheet = workbook[sheet_name]

    for row in range(2, sheet.max_row + 1):
        cell = sheet.cell(row=row, column=3)
        corrected_price = cell.value * 0.9
        print(cell.value, corrected_price)
        corrected_price_cell = sheet.cell(row=row, column=4)
        corrected_price_cell.value = corrected_price

    chart_values = Reference(
        sheet,
        min_row=2,
        max_row=sheet.max_row,
        min_col=4,
        max_col=4
    )

    chart = BarChart()
    chart.add_data(chart_values)
    sheet.add_chart(chart)

    workbook.save(output_file)
